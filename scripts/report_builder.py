# -*- coding: utf-8 -*-
"""
Module xuất báo cáo Excel đầu ra.
Tạo file Excel với nhiều sheet tương ứng từng báo cáo.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers
from openpyxl.utils import get_column_letter
from pathlib import Path
from utils import get_output_dir, get_report_month_label, format_number, safe_float


# Styles
HEADER_FONT = Font(name="Times New Roman", bold=True, size=11)
DATA_FONT = Font(name="Times New Roman", size=11)
TITLE_FONT = Font(name="Times New Roman", bold=True, size=14)
TOTAL_FONT = Font(name="Times New Roman", bold=True, size=11, color="0000FF")
WARNING_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
ERROR_FILL = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT_WHITE = Font(name="Times New Roman", bold=True, size=11, color="FFFFFF")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)
NUMBER_FORMAT = '#,##0'


def write_header_row(ws, row, headers, start_col=1):
    """Ghi dòng header với format."""
    for j, h in enumerate(headers):
        cell = ws.cell(row=row, column=start_col + j, value=h)
        cell.font = HEADER_FONT_WHITE
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def write_data_row(ws, row, values, start_col=1, is_total=False):
    """Ghi dòng dữ liệu."""
    for j, v in enumerate(values):
        cell = ws.cell(row=row, column=start_col + j)
        if isinstance(v, (int, float)) and not pd.isna(v):
            cell.value = v
            cell.number_format = NUMBER_FORMAT
            if v < 0:
                cell.font = Font(name="Times New Roman", size=11, color="FF0000",
                                bold=is_total)
            else:
                cell.font = TOTAL_FONT if is_total else DATA_FONT
        else:
            cell.value = v
            cell.font = TOTAL_FONT if is_total else DATA_FONT
        cell.border = THIN_BORDER
        cell.alignment = Alignment(vertical="center")


def write_df_to_sheet(ws, df, title=None, start_row=1):
    """Ghi DataFrame vào sheet."""
    current_row = start_row

    if title:
        ws.cell(row=current_row, column=1, value=title).font = TITLE_FONT
        current_row += 2

    # Headers
    write_header_row(ws, current_row, list(df.columns))
    current_row += 1

    # Data
    for _, row in df.iterrows():
        is_total = any(str(v).lower() in ["cộng", "tổng cộng", "tổng", "total"]
                      for v in row.values if pd.notna(v))
        values = []
        for v in row.values:
            if pd.isna(v):
                values.append("")
            else:
                values.append(v)
        write_data_row(ws, current_row, values, is_total=is_total)
        current_row += 1

    # Auto-width
    for col_idx in range(1, len(df.columns) + 1):
        max_len = max(
            len(str(df.columns[col_idx - 1])),
            *[len(str(v)) for v in df.iloc[:, col_idx - 1] if pd.notna(v)]
        ) if len(df) > 0 else len(str(df.columns[col_idx - 1]))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 40)

    return current_row


def build_report(all_results: dict, config: dict) -> str:
    """
    Tạo file Excel báo cáo đầu ra.
    Returns: đường dẫn file đã tạo.
    """
    wb = Workbook()
    month_label = get_report_month_label(config)
    report_month = config["report_month"]

    # ===== Sheet 1: TỔNG HỢP =====
    ws = wb.active
    ws.title = "TONG_HOP"
    taxvta = all_results.get("taxvta", {})

    row = 1
    ws.cell(row=row, column=1, value=f"BÁO CÁO THUẾ GTGT - {month_label}").font = TITLE_FONT
    row += 2

    # Cross-checks
    ws.cell(row=row, column=1, value="KIỂM TRA CHÉO").font = Font(bold=True, size=12)
    row += 1
    cross_checks = taxvta.get("cross_checks", [])
    if cross_checks:
        headers = ["STT", "Nội dung", "Nguồn A", "Giá trị A", "Nguồn B", "Giá trị B", "Chênh lệch", "Trạng thái"]
        write_header_row(ws, row, headers)
        row += 1
        for i, cc in enumerate(cross_checks):
            values = [i + 1, cc["name"], cc["source_a"], cc["value_a"],
                     cc["source_b"], cc["value_b"], cc["chenh_lech"], cc["status"]]
            is_warning = "LỆCH" in cc.get("status", "")
            write_data_row(ws, row, values)
            if is_warning:
                for col in range(1, len(values) + 1):
                    ws.cell(row=row, column=col).fill = WARNING_FILL
            row += 1

    # ===== Các sheet dữ liệu =====
    sheet_configs = [
        ("GCS", "bchdon", "summary", f"BC TỔNG HỢP THEO NGÀY GCS - {month_label}"),
        ("GL0903", "gl0903", "summary", f"GL_0903_TK511 - DOANH THU - {month_label}"),
        ("TA035", "ta035", "summary", f"TA_035 BẢNG KÊ BÁN RA - {month_label}"),
        ("TK33895", "gl038", "summary", f"GL_038 TK33895 - {month_label}"),
        ("Nhom_TC", "nhomtc", "summary", f"SẢN LƯỢNG NHÔM TOÀN CẦU - {month_label}"),
    ]

    for sheet_name, result_key, df_key, title in sheet_configs:
        ws_new = wb.create_sheet(sheet_name)
        result = all_results.get(result_key, {})
        df = result.get(df_key, pd.DataFrame())
        if not df.empty:
            write_df_to_sheet(ws_new, df, title=title)

    # TA036 với MST warnings
    ws_ta036 = wb.create_sheet("TA036")
    ta036 = all_results.get("ta036", {})
    ta036_summary = ta036.get("summary", pd.DataFrame())
    if not ta036_summary.empty:
        end_row = write_df_to_sheet(ws_ta036, ta036_summary,
                                     title=f"TA_036 BẢNG KÊ MUA VÀO - {month_label}")
        # MST warnings
        mst_warnings = ta036.get("mst_warnings", pd.DataFrame())
        if not mst_warnings.empty:
            end_row += 2
            ws_ta036.cell(row=end_row, column=1,
                         value=f"⚠ CẢNH BÁO MST ({len(mst_warnings)} trường hợp)").font = Font(
                         bold=True, color="FF0000", size=12)
            end_row += 1
            warn_cols = ["STT", "TenNguoiBan", "MSTNguoiBan", "MST_CHECK", "TAIKHOAN"]
            warn_df = mst_warnings[warn_cols].copy() if all(c in mst_warnings.columns for c in warn_cols) else mst_warnings
            write_df_to_sheet(ws_ta036, warn_df, start_row=end_row)

    # TA036 theo thuế suất
    ta036_ts = ta036.get("summary_thue_suat", pd.DataFrame())
    if not ta036_ts.empty:
        ws_ta036_ts = wb.create_sheet("TA036_TS")
        write_df_to_sheet(ws_ta036_ts, ta036_ts,
                         title=f"TA_036 THEO THUẾ SUẤT - {month_label}")

    # TA030_3331
    ws_3331 = wb.create_sheet("TK333111")
    ta030_3331 = all_results.get("ta030_3331", {})
    ta030_3331_sum = ta030_3331.get("summary", pd.DataFrame())
    if not ta030_3331_sum.empty:
        write_df_to_sheet(ws_3331, ta030_3331_sum,
                         title=f"TA_030 TK333111 - {month_label}")

    # TA030_1331 summary
    ws_1331 = wb.create_sheet("TK1331")
    ta030_1331 = all_results.get("ta030_1331", {})
    ta030_1331_sum = ta030_1331.get("summary", pd.DataFrame())
    if not ta030_1331_sum.empty:
        write_df_to_sheet(ws_1331, ta030_1331_sum,
                         title=f"TA_030 TK1331 - {month_label}")

    # Kiểm dò TK
    ws_kd = wb.create_sheet("KIEM_DO")
    kiem_do = taxvta.get("kiem_do_tk", {})
    row = 1
    ws_kd.cell(row=row, column=1, value=f"KIỂM DÒ THEO TÀI KHOẢN - {month_label}").font = TITLE_FONT
    row += 2

    for tk_name, tk_data in kiem_do.items():
        ws_kd.cell(row=row, column=1, value=tk_name).font = Font(bold=True, size=12)
        row += 1
        for key, val in tk_data.items():
            ws_kd.cell(row=row, column=1, value=key).font = DATA_FONT
            cell = ws_kd.cell(row=row, column=2, value=val)
            cell.font = DATA_FONT
            cell.number_format = NUMBER_FORMAT
            row += 1
        row += 1

    # Save
    output_dir = get_output_dir()
    output_dir.mkdir(parents=True, exist_ok=True)
    output_filename = config.get("output_file", "BAO_CAO_THUE_GTGT.xlsx").replace(
        "{report_month}", report_month
    )
    output_path = output_dir / output_filename
    wb.save(str(output_path))

    return str(output_path)
